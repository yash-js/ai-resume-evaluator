from pathlib import Path

from docx import Document
from pypdf import PdfReader

SUPPORTED_SUFFIXES = {".pdf", ".docx"}


def read_pdf(file_path: Path) -> str:
    reader = PdfReader(file_path)
    pages: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            pages.append(page_text)
    return "\n".join(pages)


def read_docx(file_path: Path) -> str:
    document = Document(file_path)
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells: list[str] = []
            seen: set[str] = set()
            for cell in row.cells:
                text = cell.text.strip()
                if text and text not in seen:
                    seen.add(text)
                    cells.append(text)
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def read_resume(file_path: Path) -> str | None:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return read_pdf(file_path)
    if suffix == ".docx":
        return read_docx(file_path)
    return None
